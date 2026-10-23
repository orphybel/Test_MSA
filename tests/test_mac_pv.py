"""Report des adresses MAC dans la fiche de test du NVR."""

import os
import re
import sys
import time
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
import pytest
from openpyxl.styles import Border, Side

from msa_test import mac_pv, pv_excel
from msa_test.mac_pv import CLE_SWITCH
from msa_test.pv_excel import ErreurPV, lire_cellule

TRAIT = Side(style="thin")
CADRE = Border(left=TRAIT, right=TRAIT, top=TRAIT, bottom=TRAIT)


# ---------------------------------------------------------------------- #
# Modeles de fiche reproduisant la structure X301226 / X301281
# ---------------------------------------------------------------------- #
def _fiche(chemin, modele, slots, feuille="Constit produit"):
    """Fiche avec, pour chaque module, une ligne @MAC dont F est vide."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = feuille
    ws["A11"], ws["F11"] = "Nom de l’article", "N° Série"
    ws["B12"] = modele
    ligne = 14
    blocs = [("Module Contrôle / Switch Ethernet", "Carte Contrôle / Switch")]
    blocs += [("Module CPU Enreg. Vidéo (slot n°%d)" % k, "Carte Enregistreur Vidéo") for k in range(slots)]
    for titre, carte in blocs:
        ws["A%d" % ligne] = titre
        ws["A%d" % (ligne + 1)] = carte
        ws["B%d" % (ligne + 2)] = "@MAC                                    "
        for colonne in "BCDEFG":
            ws["%s%d" % (colonne, ligne + 2)].border = CADRE  # F reste vide mais stylee
        ligne += 3
    ws["A%d" % ligne] = "Module Stockage Amovible (Slot n°0)"
    ws["B%d" % (ligne + 1)] = "@MAC"  # @MAC hors bloc CPU : ignore
    wb.save(chemin)
    return str(chemin)


@pytest.fixture
def pv_rerng(tmp_path):
    return _fiche(tmp_path / "X301281-9_RERNG_NVR_Fiche-de-Test-DUP N°2609026-0152011-00390.xlsx",
                  "RERNG-NVR-2", 2)


@pytest.fixture
def pv_mp14(tmp_path):
    return _fiche(tmp_path / "X301226_9-MP14_FicheTest_MP14-NVR-4_N°260212-0154864-00107.xlsx",
                  "MP14-NVR-4", 6, feuille="Constit produit(1_2)")


RELEVE = """RELEVE DES ADRESSES MAC
Equipement NVR - procedure X301773
==============================================================================

Date            : 2026-10-23T11:31:15
N° de serie NVR : 390
Operateur       : LG
Nombre de MSA   : 2

Carte Controle/Switch : relevée via son interface web.

RECAPITULATIF
----------------------------------------------------------------------------------------------------
EQUIPEMENT             ADRESSE IP       MODULE / INTERFACE                     ADRESSE MAC
Carte Controle/Switch  192.168.0.186    Module Contrôle Switch                 00:10:02:0f:da:ad
Carte Controle/Switch  192.168.0.186    Module CPU Enregistreur 0              00:e0:4b:8c:4b:cf
Carte Controle/Switch  192.168.0.186    Module CPU Enregistreur 1              00:e0:4b:8c:4b:ad
Carte Controle/Switch  192.168.0.186    cam1_Z11                               00:90:e8:71:f3:71
MSA0                   192.168.0.187    enp4s0                                 00:e0:4b:8c:4b:cf
MSA1                   192.168.0.188    enp4s0                                 00:e0:4b:8c:4b:ad

DETAIL PAR EQUIPEMENT
----------------------------------------------------------------------------------------------------
Carte Controle/Switch (http://192.168.0.186/) - source : interface web
  Module CPU Enregistreur 5                00:00:00:00:00:55
"""


def _releve(tmp_path, texte=RELEVE, nom="adresses_MAC_390_2026-10-23_11-31-15.txt"):
    chemin = tmp_path / nom
    chemin.write_text(texte, encoding="utf-8")
    return str(chemin)


# ---------------------------------------------------------------------- #
# Lecture du releve
# ---------------------------------------------------------------------- #
def test_lecture_du_releve(tmp_path):
    releve = mac_pv.lire_releve_mac(_releve(tmp_path))
    assert releve["serie_nvr"] == "390"
    assert releve["nombre_msa"] == 2
    assert releve["adresses"] == {
        CLE_SWITCH: "00:10:02:0f:da:ad",
        0: "00:e0:4b:8c:4b:cf",
        1: "00:e0:4b:8c:4b:ad",
    }
    assert set(releve["sources"].values()) == {"interface web"}
    assert releve["incoherences"] == []


def test_seul_le_recapitulatif_est_lu(tmp_path):
    """La section DETAIL (ici un slot 5 fictif) n'est pas prise en compte."""
    assert 5 not in mac_pv.lire_releve_mac(_releve(tmp_path))["adresses"]


def test_cameras_ignorees(tmp_path):
    adresses = mac_pv.lire_releve_mac(_releve(tmp_path))["adresses"].values()
    assert "00:90:e8:71:f3:71" not in adresses


def test_repli_sur_le_releve_ssh(tmp_path):
    texte = "\n".join(l for l in RELEVE.splitlines() if "CPU Enregistreur 1" not in l)
    releve = mac_pv.lire_releve_mac(_releve(tmp_path, texte))
    assert releve["adresses"][1] == "00:e0:4b:8c:4b:ad"
    assert releve["sources"][1] == "SSH"


def test_incoherence_web_ssh_signalee(tmp_path):
    texte = RELEVE.replace(
        "MSA1                   192.168.0.188    enp4s0                                 00:e0:4b:8c:4b:ad",
        "MSA1                   192.168.0.188    enp4s0                                 00:e0:4b:8c:4b:99",
    )
    releve = mac_pv.lire_releve_mac(_releve(tmp_path, texte))
    assert releve["adresses"][1] == "00:e0:4b:8c:4b:ad"  # la page web fait foi
    assert len(releve["incoherences"]) == 1 and "slot n°1" in releve["incoherences"][0]


def test_interfaces_virtuelles_ignorees(tmp_path):
    texte = RELEVE.replace(
        "MSA0                   192.168.0.187    enp4s0",
        "MSA0                   192.168.0.187    ifb0                                   02:00:00:00:00:01\n"
        "MSA0                   192.168.0.187    enp4s0",
    ).replace("Module CPU Enregistreur 0              00:e0:4b:8c:4b:cf\n", "")
    releve = mac_pv.lire_releve_mac(_releve(tmp_path, texte))
    assert releve["adresses"][0] == "00:e0:4b:8c:4b:cf"


def test_fichier_qui_n_est_pas_un_releve(tmp_path):
    with pytest.raises(ErreurPV):
        mac_pv.lire_releve_mac(_releve(tmp_path, "un autre fichier"))


def test_dernier_releve_du_dossier(tmp_path):
    ancien = _releve(tmp_path, nom="adresses_MAC_1.txt")
    time.sleep(0.01)
    recent = _releve(tmp_path, nom="adresses_MAC_2.txt")
    os.utime(ancien, (1, 1))
    assert mac_pv.dernier_releve_mac(str(tmp_path)) == recent


# ---------------------------------------------------------------------- #
# Reperage des cellules @MAC
# ---------------------------------------------------------------------- #
def test_cellules_du_pv_rerng(pv_rerng):
    analyse = mac_pv.analyser_pv(pv_rerng)
    assert analyse["modele"] == "RERNG-NVR"
    assert analyse["cellules"] == {CLE_SWITCH: "F16", 0: "F19", 1: "F22"}


def test_cellules_du_pv_mp14(pv_mp14):
    analyse = mac_pv.analyser_pv(pv_mp14)
    assert analyse["modele"] == "MP14-NVR"
    assert analyse["feuille"] == "Constit produit(1_2)"
    assert list(analyse["cellules"]) == [CLE_SWITCH, 0, 1, 2, 3, 4, 5]
    assert analyse["variante"] == 4


def test_pv_sans_ligne_mac(tmp_path):
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "autre.xlsx")
    with pytest.raises(ErreurPV, match="@MAC"):
        mac_pv.analyser_pv(str(tmp_path / "autre.xlsx"))


@pytest.mark.parametrize(
    "nom,attendu",
    [
        ("X301226_9-MP14_FicheTest_MP14-NVR-4_N°260212-0154864-00107.xlsx", 4),
        ("X301226_9-MP14_FicheTest_MP14-NVR-X_N°260408-0147460-00196.xlsx", "X"),
        ("X301226_8-MP14_FicheTest_MP14-NVR-6 N° 250411-0135006-00057.xlsx", 6),
        ("X301281-9_RERNG_NVR_Fiche-de-Test-DUP N°2609026-0152011-00385.xlsm", None),
    ],
)
def test_variante_d_apres_le_nom(nom, attendu):
    assert mac_pv.variante_du_nom(nom) == attendu


# ---------------------------------------------------------------------- #
# PV concordant
# ---------------------------------------------------------------------- #
NOMS = [
    "X301226_9-MP14_FicheTest_MP14-NVR-4_N°260212-0154864-00107.xlsx",
    "X301226_9-MP14_FicheTest_MP14-NVR-X_N°260408-0147460-00196.xlsx",
    "X301226_8-MP14_FicheTest_MP14-NVR-6 N° 250411-0135006-00057.xlsx",
    "X301281-9_RERNG_NVR_Fiche-de-Test-DUP N°2609026-0152011-00385.xlsm",
    "~$X301281-9_RERNG_NVR_Fiche-de-Test-DUP N°2609026-0152011-00385.xlsm",
    "notes.txt",
]


@pytest.fixture
def dossier_pv(tmp_path):
    for nom in NOMS:
        (tmp_path / nom).write_bytes(b"x")
    return str(tmp_path)


@pytest.mark.parametrize(
    "serie,attendu",
    [
        ("107", NOMS[0]),
        ("196", NOMS[1]),
        ("57", NOMS[2]),  # 00057
        ("385", NOMS[3]),  # le fichier temporaire ~$ est ignore
        ("2609026-0152011-00385", NOMS[3]),
    ],
)
def test_pv_concordant(dossier_pv, serie, attendu):
    assert [os.path.basename(p) for p in mac_pv.pv_concordants(dossier_pv, serie)] == [attendu]


def test_aucun_pv_concordant(dossier_pv):
    assert mac_pv.pv_concordants(dossier_pv, "999") == []
    assert mac_pv.pv_concordants(dossier_pv, "") == []
    assert mac_pv.pv_concordants("/dossier/absent", "107") == []


# ---------------------------------------------------------------------- #
# Preparation et ecriture
# ---------------------------------------------------------------------- #
def test_plan_d_ecriture(pv_rerng, tmp_path):
    plan = mac_pv.preparer(pv_rerng, mac_pv.lire_releve_mac(_releve(tmp_path)))
    assert plan["erreurs"] == [] and plan["alertes"] == []
    assert [(l["cellule"], l["mac"]) for l in plan["lignes"]] == [
        ("F16", "00:10:02:0F:DA:AD"),
        ("F19", "00:E0:4B:8C:4B:CF"),
        ("F22", "00:E0:4B:8C:4B:AD"),
    ]


def test_ecriture_et_conservation_du_classeur(pv_rerng, tmp_path):
    releve = mac_pv.lire_releve_mac(_releve(tmp_path))
    with zipfile.ZipFile(pv_rerng) as archive:
        avant = {n: archive.read(n) for n in archive.namelist()}
    sauvegarde = mac_pv.ecrire(mac_pv.preparer(pv_rerng, releve), str(tmp_path / "sauv"))

    assert lire_cellule(pv_rerng, "Constit produit", "F16") == "00:10:02:0F:DA:AD"
    assert lire_cellule(pv_rerng, "Constit produit", "F19") == "00:E0:4B:8C:4B:CF"
    assert lire_cellule(pv_rerng, "Constit produit", "F22") == "00:E0:4B:8C:4B:AD"

    with zipfile.ZipFile(pv_rerng) as archive:
        apres = {n: archive.read(n) for n in archive.namelist()}
    assert list(avant) == list(apres)
    assert {n for n in avant if avant[n] != apres[n]} == {"xl/worksheets/sheet1.xml"}

    # aucune autre cellule ne disparait (cellules F vides auto-fermantes)
    cellules = lambda xml: re.findall(r'<c r="([A-Z]+\d+)"', xml.decode())
    assert cellules(avant["xl/worksheets/sheet1.xml"]) == cellules(apres["xl/worksheets/sheet1.xml"])
    ws = openpyxl.load_workbook(pv_rerng)["Constit produit"]
    assert ws["A17"].value == "Module CPU Enreg. Vidéo (slot n°0)"
    assert ws["G16"].border.left.style == "thin"

    # copie du PV d'origine
    assert os.path.basename(sauvegarde).startswith(
        "X301281-9_RERNG_NVR_Fiche-de-Test-DUP N°2609026-0152011-00390_avant_MAC_"
    )
    with zipfile.ZipFile(sauvegarde) as archive:
        assert {n: archive.read(n) for n in archive.namelist()} == avant


def test_seconde_ecriture_sans_remplacement(pv_rerng, tmp_path):
    releve = mac_pv.lire_releve_mac(_releve(tmp_path))
    mac_pv.ecrire(mac_pv.preparer(pv_rerng, releve))
    plan = mac_pv.preparer(pv_rerng, releve)
    assert [l["actuelle"] for l in plan["lignes"]] == [l["mac"] for l in plan["lignes"]]
    assert mac_pv.remplacements(plan) == []


def test_adresse_differente_signalee(pv_rerng, tmp_path):
    releve = mac_pv.lire_releve_mac(_releve(tmp_path))
    mac_pv.ecrire(mac_pv.preparer(pv_rerng, releve))
    releve["adresses"][0] = "aa:bb:cc:dd:ee:ff"
    remplacees = mac_pv.remplacements(mac_pv.preparer(pv_rerng, releve))
    assert [(l["cellule"], l["actuelle"], l["mac"]) for l in remplacees] == [
        ("F19", "00:E0:4B:8C:4B:CF", "AA:BB:CC:DD:EE:FF")
    ]


def test_slots_non_equipes_laisses_vides(pv_mp14, tmp_path):
    releve = mac_pv.lire_releve_mac(_releve(tmp_path))
    plan = mac_pv.preparer(pv_mp14, releve)
    assert [l["cle"] for l in plan["lignes"]] == [CLE_SWITCH, 0, 1]
    mac_pv.ecrire(plan)
    assert lire_cellule(pv_mp14, "Constit produit(1_2)", "F25") in (None, "")


def test_trop_de_modules_pour_le_pv(pv_rerng, tmp_path):
    releve = mac_pv.lire_releve_mac(_releve(tmp_path))
    releve["adresses"].update({k: "00:00:00:00:00:%02x" % k for k in range(2, 6)})
    plan = mac_pv.preparer(pv_rerng, releve)
    assert plan["erreurs"] and "ce n'est pas le PV de ce NVR" in plan["erreurs"][0]
    with pytest.raises(ErreurPV):
        mac_pv.ecrire(plan)


def test_variante_incoherente_signalee(pv_mp14, tmp_path):
    """Le nom indique MP14-NVR-4 mais seulement 2 modules ont ete releves."""
    plan = mac_pv.preparer(pv_mp14, mac_pv.lire_releve_mac(_releve(tmp_path)))
    assert any("MP14-NVR-4" in a for a in plan["alertes"])
    assert plan["erreurs"] == []


def _fichiers_temporaires(dossier):
    return [f for f in os.listdir(dossier) if f.endswith(".tmp")]


def test_pv_non_tenu_ouvert_pendant_le_remplacement(pv_rerng, tmp_path, monkeypatch):
    """Sous Windows, un fichier ouvert (meme par le logiciel) ne peut pas etre
    remplace : [WinError 5] Acces refuse sur le partage reseau."""
    plan = mac_pv.preparer(pv_rerng, mac_pv.lire_releve_mac(_releve(tmp_path)))
    remplacement = os.replace
    ouverts = []

    def verifier(source, destination):
        cible = os.path.realpath(destination)
        for fd in os.listdir("/proc/self/fd"):
            try:
                if os.path.realpath(os.readlink("/proc/self/fd/" + fd)) == cible:
                    ouverts.append(fd)
            except OSError:
                pass
        remplacement(source, destination)

    monkeypatch.setattr("msa_test.pv_excel.os.replace", verifier)
    mac_pv.ecrire(plan)
    assert ouverts == []


def test_renommage_refuse_sur_le_reseau(pv_rerng, tmp_path, monkeypatch):
    """Partage autorisant la modification mais pas le renommage : le contenu
    est recopie par-dessus le PV, sans fichier temporaire residuel."""
    plan = mac_pv.preparer(pv_rerng, mac_pv.lire_releve_mac(_releve(tmp_path)))

    def refuse(source, destination):
        erreur = PermissionError(13, "Acces refuse")
        erreur.winerror = 5
        raise erreur

    monkeypatch.setattr("msa_test.pv_excel.os.replace", refuse)
    monkeypatch.setattr("msa_test.pv_excel._ATTENTE", 0)
    mac_pv.ecrire(plan)
    relu = mac_pv.analyser_pv(pv_rerng)
    assert all(relu["valeurs"][l["cellule"]] == l["mac"] for l in plan["lignes"])
    assert _fichiers_temporaires(os.path.dirname(pv_rerng)) == []


def test_pv_ouvert_dans_excel(pv_rerng, tmp_path, monkeypatch):
    """Ni renommage ni ecriture possibles : message clair, PV intact."""
    plan = mac_pv.preparer(pv_rerng, mac_pv.lire_releve_mac(_releve(tmp_path)))
    avant = open(pv_rerng, "rb").read()

    def verrouille(*args):
        erreur = PermissionError(13, "Le processus ne peut pas acceder au fichier")
        erreur.winerror = 32
        raise erreur

    monkeypatch.setattr("msa_test.pv_excel.os.replace", verrouille)
    monkeypatch.setattr("msa_test.pv_excel.shutil.copyfile", verrouille)
    monkeypatch.setattr("msa_test.pv_excel._ATTENTE", 0)
    with pytest.raises(ErreurPV, match="utilise par un autre programme"):
        mac_pv.ecrire(plan)
    assert open(pv_rerng, "rb").read() == avant
    assert _fichiers_temporaires(os.path.dirname(pv_rerng)) == []


def test_dossier_sans_droit_de_creation(pv_rerng, tmp_path, monkeypatch):
    """Creation de fichier interdite dans le dossier du PV : le fichier
    temporaire est construit dans le dossier temporaire de Windows."""
    plan = mac_pv.preparer(pv_rerng, mac_pv.lire_releve_mac(_releve(tmp_path)))
    dossier_pv = os.path.dirname(os.path.abspath(pv_rerng))
    construire = pv_excel._construire_archive
    autre = tmp_path / "temp_windows"
    autre.mkdir()

    def construction(source, chemin, modifications):
        if os.path.dirname(os.path.abspath(chemin)) == dossier_pv:
            raise PermissionError(13, "Acces refuse")
        construire(source, chemin, modifications)

    monkeypatch.setattr(pv_excel, "_construire_archive", construction)
    monkeypatch.setattr(pv_excel.tempfile, "gettempdir", lambda: str(autre))
    mac_pv.ecrire(plan)
    relu = mac_pv.analyser_pv(pv_rerng)
    assert all(relu["valeurs"][l["cellule"]] == l["mac"] for l in plan["lignes"])
    assert os.listdir(autre) == []


def test_sauvegarde_impossible_pv_intact(pv_rerng, tmp_path, monkeypatch):
    plan = mac_pv.preparer(pv_rerng, mac_pv.lire_releve_mac(_releve(tmp_path)))
    avant = open(pv_rerng, "rb").read()
    bloque = tmp_path / "fichier"
    bloque.write_text("x")
    with pytest.raises(ErreurPV, match="Sauvegarde impossible.*pas ete modifie"):
        mac_pv.ecrire(plan, str(bloque / "sauvegardes"))
    assert open(pv_rerng, "rb").read() == avant


def test_chemins_longs_windows(monkeypatch):
    monkeypatch.setattr(pv_excel.sys, "platform", "win32")
    monkeypatch.setattr(pv_excel.os.path, "abspath", lambda p: p)
    court = "C:\\PV\\a.xlsm"
    assert pv_excel.chemin_systeme(court) == court
    long_local = "C:\\" + "d" * 300 + "\\a.xlsm"
    assert pv_excel.chemin_systeme(long_local) == "\\\\?\\" + long_local
    long_unc = "\\\\srvfich01\\DUP\\" + "d" * 300 + "\\a.xlsm"
    assert pv_excel.chemin_systeme(long_unc) == (
        "\\\\?\\UNC\\srvfich01\\DUP\\" + "d" * 300 + "\\a.xlsm")


# ---------------------------------------------------------------------- #
# PV complet : les autres onglets ne doivent pas etre modifies
# ---------------------------------------------------------------------- #
def _pv_multi_onglets(chemin, piege_constit=False):
    """Fiche a plusieurs onglets, comme les PV reels."""
    _fiche(chemin, "MP14-NVR-6", 6, feuille="Constit produit(1_2)")
    wb = openpyxl.load_workbook(chemin)
    suite = wb.create_sheet("Constit produit(2_2)")
    suite["A12"] = "Module Stockage Amovible (Slot n°0)"
    final = wb.create_sheet("Test final (1_1)")
    final["B2"] = "='Constit produit(1_2)'!F16"
    final["A10"], final["C10"] = "Mesure", 12.5
    mesures = wb.create_sheet("App Mesures")
    mesures.sheet_state = "hidden"
    # piege : un onglet de mesures qui contient lui aussi un bloc et une ligne @MAC
    mesures["A3"] = "Module Contrôle / Switch Ethernet"
    mesures["B5"] = "@MAC"
    mesures["F5"] = "ne pas toucher"
    if piege_constit:
        suite["A14"] = "Module Contrôle / Switch Ethernet"
        suite["B16"] = "@MAC"
    wb.save(chemin)
    return str(chemin)


def test_seul_l_onglet_de_constitution_est_modifie(tmp_path):
    pv = _pv_multi_onglets(tmp_path / "X301226_9-MP14_FicheTest_MP14-NVR-6 N°250411-0135006-00057.xlsx")
    with zipfile.ZipFile(pv) as archive:
        avant = {n: archive.read(n) for n in archive.namelist()}
        feuilles = mac_pv._chemins_feuilles(archive)

    plan = mac_pv.preparer(pv, mac_pv.lire_releve_mac(_releve(tmp_path)))
    assert plan["pv"]["feuille"] == "Constit produit(1_2)"
    assert any("App Mesures" in a for a in plan["alertes"])  # onglet ecarte signale
    mac_pv.ecrire(plan)

    with zipfile.ZipFile(pv) as archive:
        apres = {n: archive.read(n) for n in archive.namelist()}
    modifiees = {n for n in avant if avant[n] != apres[n]}
    assert modifiees == {feuilles["Constit produit(1_2)"]}
    # chaque autre onglet, le classeur, les styles et les textes sont identiques
    for nom in ("Constit produit(2_2)", "Test final (1_1)", "App Mesures"):
        assert apres[feuilles[nom]] == avant[feuilles[nom]], nom
    communs = [e for e in ("xl/workbook.xml", "xl/styles.xml", "xl/sharedStrings.xml") if e in avant]
    assert "xl/workbook.xml" in communs and "xl/styles.xml" in communs
    for entree in communs:
        assert apres[entree] == avant[entree], entree

    wb = openpyxl.load_workbook(pv)
    assert wb["App Mesures"].sheet_state == "hidden"
    assert wb["App Mesures"]["F5"].value == "ne pas toucher"
    assert wb["Test final (1_1)"]["B2"].value == "='Constit produit(1_2)'!F16"
    assert wb["Constit produit(1_2)"]["F16"].value == "00:10:02:0F:DA:AD"


def test_plusieurs_onglets_de_constitution_possibles(tmp_path):
    """Deux onglets "Constit produit" avec des lignes @MAC : refus, pas de devinette."""
    pv = _pv_multi_onglets(tmp_path / "pv.xlsx", piege_constit=True)
    avant = open(pv, "rb").read()
    with pytest.raises(ErreurPV, match="Plusieurs onglets"):
        mac_pv.preparer(pv, mac_pv.lire_releve_mac(_releve(tmp_path)))
    assert open(pv, "rb").read() == avant


# ---------------------------------------------------------------------- #
# Traitement par lot : dossier de releves + dossier de PV
# ---------------------------------------------------------------------- #
def _releve_nvr(dossier, serie, date, cpu0="00:e0:4b:8c:4b:cf", nom=None):
    texte = RELEVE.replace("N° de serie NVR : 390", "N° de serie NVR : %s" % serie)
    texte = texte.replace("2026-10-23T11:31:15", date).replace("00:e0:4b:8c:4b:cf", cpu0)
    os.makedirs(dossier, exist_ok=True)
    chemin = os.path.join(dossier, nom or "adresses_MAC_%s_%s.txt" % (serie, date.replace(":", "-")))
    with open(chemin, "w", encoding="utf-8") as fichier:
        fichier.write(texte)
    return chemin


@pytest.fixture
def lot(tmp_path):
    releves, pvs = tmp_path / "releves", tmp_path / "pv"
    pvs.mkdir()
    _fiche(pvs / "X301281-9_RERNG_NVR_Fiche-de-Test-DUP N°2609026-0152011-00390.xlsx", "RERNG-NVR-2", 2)
    _fiche(pvs / "X301281-9_RERNG_NVR_Fiche-de-Test-DUP N°2609026-0152011-00391.xlsx", "RERNG-NVR-2", 2)
    _fiche(pvs / "X301226_8-MP14_FicheTest_MP14-NVR-X_N°260408-0147460-00196.xlsx", "MP14-NVR-X", 6)
    _fiche(pvs / "X301226_9-MP14_FicheTest_MP14-NVR-X_N°260408-0147460-00196.xlsx", "MP14-NVR-X", 6)
    _releve_nvr(str(releves), "390", "2026-10-20T09:00:00", cpu0="00:e0:4b:00:00:01")  # ancien
    _releve_nvr(str(releves), "390", "2026-10-23T11:31:15")
    _releve_nvr(str(releves / "jour2"), "391", "2026-10-22T10:00:00")  # sous-dossier
    _releve_nvr(str(releves), "196", "2026-10-22T10:00:00")
    _releve_nvr(str(releves), "999", "2026-10-22T10:00:00")
    _releve_nvr(str(releves), "...", "2026-10-22T12:00:00", nom="adresses_MAC_sans_numero.txt")
    (releves / "adresses_MAC_casse.txt").write_text("pas un releve", encoding="utf-8")
    (releves / "notes.txt").write_text("ignore", encoding="utf-8")
    return str(releves), str(pvs)


def _par_nvr(entrees):
    return {(e["serie"], os.path.basename(e["chemin_releve"])): e for e in entrees}


def test_releves_du_dossier_et_sous_dossiers(lot):
    noms = [os.path.basename(c) for c in mac_pv.releves_du_dossier(lot[0])]
    assert len(noms) == 7  # notes.txt ignore, sous-dossier inclus
    assert "notes.txt" not in noms


def test_etats_du_lot(lot):
    entrees = mac_pv.preparer_lot(*lot)
    etats = {(e["serie"] or os.path.basename(e["chemin_releve"])): e["etat"] for e in entrees
             if e["etat"] != mac_pv.IGNORE}
    assert etats["390"] == mac_pv.PRET
    assert etats["391"] == mac_pv.PRET
    assert etats["196"] == mac_pv.PLUSIEURS_PV
    assert etats["999"] == mac_pv.SANS_PV
    assert etats["adresses_MAC_sans_numero.txt"] == mac_pv.SANS_NUMERO
    assert etats["adresses_MAC_casse.txt"] == mac_pv.ERREUR


def test_le_releve_le_plus_recent_l_emporte(lot):
    entrees = [e for e in mac_pv.preparer_lot(*lot) if e["serie"] == "390"]
    retenu = [e for e in entrees if e["etat"] != mac_pv.IGNORE]
    ignore = [e for e in entrees if e["etat"] == mac_pv.IGNORE]
    assert len(retenu) == 1 and len(ignore) == 1
    assert "2026-10-23" in retenu[0]["chemin_releve"]
    assert retenu[0]["plan"]["lignes"][1]["mac"] == "00:E0:4B:8C:4B:CF"
    assert os.path.basename(retenu[0]["chemin_releve"]) in ignore[0]["message"]


def test_chaque_nvr_trouve_son_pv(lot):
    entrees = {e["serie"]: e for e in mac_pv.preparer_lot(*lot) if e["etat"] == mac_pv.PRET}
    assert entrees["390"]["pv"].endswith("-00390.xlsx")
    assert entrees["391"]["pv"].endswith("-00391.xlsx")


def test_plusieurs_pv_sans_devinette_puis_association(lot):
    entree = [e for e in mac_pv.preparer_lot(*lot) if e["serie"] == "196"][0]
    assert entree["pv"] is None and len(entree["pvs_possibles"]) == 2
    mac_pv.preparer_entree(entree, entree["pvs_possibles"][0])
    assert entree["etat"] == mac_pv.PRET


def test_ecriture_du_lot(lot):
    entrees = mac_pv.preparer_lot(*lot)
    resultats = mac_pv.ecrire_lot(entrees, os.path.join(lot[1], "..", "sauv"))
    assert sorted(e["serie"] for e, _, erreur in resultats if erreur is None) == ["390", "391"]
    pv390 = [e["pv"] for e in entrees if e["serie"] == "390" and e["pv"]][0]
    assert lire_cellule(pv390, "Constit produit", "F19") == "00:E0:4B:8C:4B:CF"
    # deuxieme analyse : tout est a jour, rien a reecrire
    apres = mac_pv.preparer_lot(*lot)
    assert {e["serie"]: e["etat"] for e in apres if e["serie"] in ("390", "391") and e["etat"] != mac_pv.IGNORE} \
        == {"390": mac_pv.A_JOUR, "391": mac_pv.A_JOUR}
    assert mac_pv.ecrire_lot(apres) == []


def test_alerte_et_remplacement_demandent_confirmation(lot, tmp_path):
    pvs = lot[1]
    _fiche(os.path.join(pvs, "X301226_9-MP14_FicheTest_MP14-NVR-4_N°260212-0154864-00107.xlsx"),
           "MP14-NVR-4", 6)
    _releve_nvr(lot[0], "107", "2026-10-22T10:00:00")  # 2 CPU pour un NVR-4 : alerte
    entrees = mac_pv.preparer_lot(*lot)
    e107 = [e for e in entrees if e["serie"] == "107"][0]
    assert e107["etat"] == mac_pv.A_VERIFIER

    sans_accord = mac_pv.ecrire_lot(entrees)
    assert "107" not in [e["serie"] for e, _, _ in sans_accord]
    avec_accord = mac_pv.ecrire_lot(mac_pv.preparer_lot(*lot), avec_confirmation=True)
    assert "107" in [e["serie"] for e, _, erreur in avec_accord if erreur is None]


def test_un_echec_n_arrete_pas_le_lot(lot, monkeypatch):
    entrees = mac_pv.preparer_lot(*lot)
    vrai_ecrire = mac_pv.ecrire

    def ecrire(plan, dossier=None):
        if plan["pv"]["chemin"].endswith("-00390.xlsx"):
            raise ErreurPV("PV ouvert dans Excel")
        return vrai_ecrire(plan, dossier)

    monkeypatch.setattr(mac_pv, "ecrire", ecrire)
    resultats = {e["serie"]: erreur for e, _, erreur in mac_pv.ecrire_lot(entrees)}
    assert resultats["390"] == "PV ouvert dans Excel"
    assert resultats["391"] is None


def test_dossier_vide_ou_absent(tmp_path):
    assert mac_pv.preparer_lot(str(tmp_path), str(tmp_path)) == []
    assert mac_pv.preparer_lot(str(tmp_path / "absent"), "") == []
