"""Remplissage des fiches de test Excel (PV)."""

import datetime
import os
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
import pytest

from msa_test import rapport
from msa_test.pv_excel import (
    ErreurPV,
    FEUILLE_CONSTITUTION,
    FEUILLE_TEST_FINAL,
    generer_pvs,
    lire_cellule,
    nom_du_pv,
    remplir_pv,
    valeurs_du_module,
)

ANCIEN_NUMERO = "0051137_0006"


@pytest.fixture
def modele(tmp_path):
    """Classeur reprenant la structure des fiches X255082-1 (cellules utiles)."""
    wb = openpyxl.Workbook()
    wb.active.title = "DENOM PRODUIT"
    wb["DENOM PRODUIT"].sheet_state = "hidden"

    constitution = wb.create_sheet(FEUILLE_CONSTITUTION)
    constitution["C2"] = "='Constit produit'!F12"
    constitution["A12"] = "Module Stockage Amovible"
    constitution["F12"] = ANCIEN_NUMERO
    constitution["A23"] = "Date:"
    constitution["B23"] = datetime.datetime(2026, 1, 15)
    constitution["B23"].number_format = "dd/mm/yyyy"
    constitution["A24"] = "Par:"
    constitution["B24"] = "QD"

    final = wb.create_sheet(FEUILLE_TEST_FINAL)
    final["B2"] = "='Constit produit'!C2"
    final["A20"] = 'Donnée SMART "UDMA_CRC_Error_Count" du disque 1 du MSA'
    final["C20"] = 0
    final["D20"] = "PASS"
    final["A21"] = 'Donnée SMART "UDMA_CRC_Error_Count" du disque 2 du MSA'
    final["C21"] = 0
    final["D21"] = "PASS"

    chemin = tmp_path / ("X255082-1_MP14_MSA_Fiche-de-Test-SAV_%s.xlsx" % ANCIEN_NUMERO)
    wb.save(chemin)
    return str(chemin)


VALEURS = {1: ("0", "PASS"), 2: ("12", "FAIL")}
DATE = datetime.date(2026, 9, 24)


def _relire(chemin):
    wb = openpyxl.load_workbook(chemin)
    return wb[FEUILLE_CONSTITUTION], wb[FEUILLE_TEST_FINAL]


# ---------------------------------------------------------------------- #
# Cellules renseignees
# ---------------------------------------------------------------------- #
def test_cellules_de_la_fiche(modele, tmp_path):
    sortie = str(tmp_path / "pv.xlsx")
    remplir_pv(modele, sortie, "0051200_0001", "J. DURAND", VALEURS, DATE)
    constitution, final = _relire(sortie)

    assert constitution["F12"].value == "0051200_0001"
    assert constitution["B23"].value == datetime.datetime(2026, 9, 24)
    assert constitution["B24"].value == "J. DURAND"
    assert final["C20"].value == 0
    assert final["C21"].value == 12
    assert final["D20"].value == "PASS"
    assert final["D21"].value == "FAIL"


def test_la_date_garde_son_format(modele, tmp_path):
    """B23 reste une date : le style de la cellule est conserve."""
    sortie = str(tmp_path / "pv.xlsx")
    remplir_pv(modele, sortie, "N1", "OP", VALEURS, DATE)
    constitution, _ = _relire(sortie)
    assert constitution["B23"].number_format == "dd/mm/yyyy"
    assert constitution["B23"].is_date


def test_le_numero_est_repris_dans_les_en_tetes(modele, tmp_path):
    """Les formules qui affichent le numero voient leur valeur actualisee."""
    sortie = str(tmp_path / "pv.xlsx")
    remplir_pv(modele, sortie, "0051200_0001", "OP", VALEURS, DATE)
    assert lire_cellule(sortie, FEUILLE_CONSTITUTION, "C2") == "0051200_0001"
    assert lire_cellule(sortie, FEUILLE_TEST_FINAL, "B2") == "0051200_0001"
    # la formule elle-meme est conservee
    constitution, final = _relire(sortie)
    assert constitution["C2"].value == "='Constit produit'!F12"
    assert final["B2"].value == "='Constit produit'!C2"


def test_excel_recalcule_a_l_ouverture(modele, tmp_path):
    sortie = str(tmp_path / "pv.xlsx")
    remplir_pv(modele, sortie, "N1", "OP", VALEURS, DATE)
    with zipfile.ZipFile(sortie) as archive:
        assert 'fullCalcOnLoad="1"' in archive.read("xl/workbook.xml").decode()


def test_valeur_non_numerique_ecrite_en_texte(modele, tmp_path):
    sortie = str(tmp_path / "pv.xlsx")
    remplir_pv(modele, sortie, "N1", "OP", {1: ("0 0 0", "PASS"), 2: (None, "FAIL")}, DATE)
    _, final = _relire(sortie)
    assert final["C20"].value == "0 0 0"
    assert final["C21"].value in (None, "")


def test_caracteres_speciaux_echappes(modele, tmp_path):
    sortie = str(tmp_path / "pv.xlsx")
    remplir_pv(modele, sortie, "A&B<1>", 'O"P', VALEURS, DATE)
    constitution, _ = _relire(sortie)
    assert constitution["F12"].value == "A&B<1>"
    assert constitution["B24"].value == 'O"P'


# ---------------------------------------------------------------------- #
# Conservation du classeur
# ---------------------------------------------------------------------- #
def test_le_reste_du_classeur_est_recopie_a_l_identique(modele, tmp_path):
    sortie = str(tmp_path / "pv.xlsx")
    remplir_pv(modele, sortie, "N1", "OP", VALEURS, DATE)
    with zipfile.ZipFile(modele) as a, zipfile.ZipFile(sortie) as b:
        assert a.namelist() == b.namelist()
        modifiees = {n for n in a.namelist() if a.read(n) != b.read(n)}
    # seules les deux feuilles renseignees et workbook.xml changent
    assert modifiees <= {
        "xl/workbook.xml",
        "xl/worksheets/sheet2.xml",
        "xl/worksheets/sheet3.xml",
    }
    assert "xl/styles.xml" not in modifiees


def test_la_feuille_cachee_reste_cachee(modele, tmp_path):
    sortie = str(tmp_path / "pv.xlsx")
    remplir_pv(modele, sortie, "N1", "OP", VALEURS, DATE)
    wb = openpyxl.load_workbook(sortie)
    assert wb["DENOM PRODUIT"].sheet_state == "hidden"


def test_le_modele_n_est_pas_modifie(modele, tmp_path):
    avant = open(modele, "rb").read()
    remplir_pv(modele, str(tmp_path / "pv.xlsx"), "N1", "OP", VALEURS, DATE)
    assert open(modele, "rb").read() == avant


# ---------------------------------------------------------------------- #
# Erreurs
# ---------------------------------------------------------------------- #
def test_feuille_absente(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Autre"
    chemin = str(tmp_path / "mauvais.xlsx")
    wb.save(chemin)
    with pytest.raises(ErreurPV):
        remplir_pv(chemin, str(tmp_path / "pv.xlsx"), "N1", "OP", VALEURS, DATE)


def test_modele_illisible(tmp_path):
    chemin = tmp_path / "pas_un_classeur.xlsm"
    chemin.write_text("texte")
    with pytest.raises(ErreurPV):
        remplir_pv(str(chemin), str(tmp_path / "pv.xlsm"), "N1", "OP", VALEURS, DATE)


def test_numero_vide(modele, tmp_path):
    with pytest.raises(ErreurPV):
        remplir_pv(modele, str(tmp_path / "pv.xlsx"), "  ", "OP", VALEURS, DATE)


# ---------------------------------------------------------------------- #
# Nom du fichier produit
# ---------------------------------------------------------------------- #
def test_nom_reprend_le_modele_avec_le_nouveau_numero(modele):
    assert (
        nom_du_pv(modele, "0051200_0001")
        == "X255082-1_MP14_MSA_Fiche-de-Test-SAV_0051200_0001.xlsx"
    )


def test_nom_quand_le_modele_ne_porte_pas_le_numero(tmp_path, modele):
    autre = tmp_path / "Fiche_vierge.xlsx"
    autre.write_bytes(open(modele, "rb").read())
    assert nom_du_pv(str(autre), "N1") == "Fiche_vierge_N1.xlsx"


def test_nom_sans_caracteres_interdits(modele):
    assert "/" not in nom_du_pv(modele, "0051/200")


# ---------------------------------------------------------------------- #
# Sanction et generation par module
# ---------------------------------------------------------------------- #
def _module(msa, sda, sdb, erreur=None):
    return {
        "msa": msa,
        "ip": "192.168.0.%d" % (187 + msa),
        "erreur": erreur,
        "partitions": {}
        if erreur
        else {
            "/dev/sda1": {"udma_crc_error_count": sda},
            "/dev/sdb1": {"udma_crc_error_count": sdb},
        },
    }


def test_sanction_par_disque():
    valeurs = valeurs_du_module(_module(0, "0", "15"), _module(0, "0", "12"))
    assert valeurs == {1: ("0", "PASS"), 2: ("15", "FAIL")}


def test_sanction_sans_releve_avant():
    valeurs = valeurs_du_module(_module(0, "0", "0"), None)
    assert valeurs[1][1] == "FAIL" and valeurs[2][1] == "FAIL"


def test_un_pv_par_msa_numerote(modele, tmp_path):
    avant = {"modules": [_module(0, "0", "0"), _module(1, "0", "0"), _module(2, "0", "0")]}
    apres = {
        "modules": [
            _module(0, "0", "0"),
            _module(1, "0", "0"),
            _module(2, None, None, erreur="SSH KO"),
        ]
    }
    produits, anomalies = generer_pvs(
        modele,
        str(tmp_path / "sortie"),
        {0: "0051200_0001", 1: "", 2: "0051200_0003"},
        avant,
        apres,
        "OP",
        date=DATE,
    )
    assert [os.path.basename(p) for p in produits] == [
        "X255082-1_MP14_MSA_Fiche-de-Test-SAV_0051200_0001.xlsx"
    ]
    assert any("MSA1" in a and "numero" in a for a in anomalies)
    assert any("MSA2" in a and "apres" in a for a in anomalies)


def test_les_fail_sont_signales(modele, tmp_path):
    avant = {"modules": [_module(0, "0", "3")]}
    apres = {"modules": [_module(0, "0", "4")]}
    produits, anomalies = generer_pvs(
        modele, str(tmp_path), {0: "N1"}, avant, apres, "OP", date=DATE
    )
    assert len(produits) == 1
    assert any("disque 2 en FAIL" in a for a in anomalies)
    _, final = _relire(produits[0])
    assert final["D21"].value == "FAIL"


# ---------------------------------------------------------------------- #
# Dossier d'enregistrement
# ---------------------------------------------------------------------- #
def test_dossier_d_enregistrement_choisi(tmp_path):
    choisi = tmp_path / "PV clients"
    try:
        rapport.choisir_dossier(str(choisi))
        assert rapport.dossier_resultats("/ailleurs") == str(choisi)
        assert choisi.is_dir()
    finally:
        rapport.choisir_dossier(None)
    assert rapport.dossier_resultats(str(tmp_path)).endswith("resultats_msa")
